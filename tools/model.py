"""Đọc 'dư liệu hôm nay.xlsx' → mô hình lô / cont chuẩn hoá (theo luật nạp 22/09, mốc ngày mới)."""
import openpyxl, datetime, re, json, sys
from collections import OrderedDict

VN = datetime.timezone(datetime.timedelta(hours=7))
def s(v):
    if v is None: return ''
    if isinstance(v, float) and v.is_integer(): v = int(v)
    return str(v).strip()
def norm(v): return re.sub(r'\s+', ' ', s(v)).strip()
def d(v):
    """date"""
    if isinstance(v, datetime.datetime): return v.date()
    if isinstance(v, datetime.date): return v
    t = s(v)
    m = re.match(r'^(\d{1,2})/(\d{1,2})/(\d{2,4})$', t)
    if m:
        y = int(m.group(3)); y = y + 2000 if y < 100 else y
        return datetime.date(y, int(m.group(2)), int(m.group(1)))
    return None
def ts(v):
    """timestamp (VN)"""
    if isinstance(v, datetime.datetime): return v.replace(tzinfo=VN)
    t = s(v)
    m = re.search(r'(\d{1,2}):(\d{2})\s*\|?\s*(\d{1,2})/+(\d{1,2})/+(\d{4})', t)   # '14:30 25/09/2026' hoặc '22:00 | 01/02/2026 ( CÁT LÁI )'
    if m:
        return datetime.datetime(int(m.group(5)), int(m.group(4)), int(m.group(3)), int(m.group(1)), int(m.group(2)), tzinfo=VN)
    m = re.search(r'(\d{1,2})/(\d{1,2})/(\d{4})\s+(\d{1,2}):(\d{2})', t)
    if m:
        return datetime.datetime(int(m.group(3)), int(m.group(2)), int(m.group(1)), int(m.group(4)), int(m.group(5)), tzinfo=VN)
    return None
def is_cont(x): return bool(re.match(r'^[A-Z]{4}\d{7}$', norm(x).upper()))

def merged_map(ws):
    mm = {}
    for rg in ws.merged_cells.ranges:
        v = ws.cell(rg.min_row, rg.min_col).value
        for r in range(rg.min_row, rg.max_row + 1):
            for c in range(rg.min_col, rg.max_col + 1): mm[(r, c)] = v
    return mm

def read_dieu_do(wb):
    ws = wb['MỚI ĐIỀU ĐỘ']; mm = merged_map(ws)
    val = lambda r, c: mm.get((r, c), ws.cell(r, c).value)
    cols = ['lo','booking','cont','seal','ma_don','ngay_goi','ngay_len_kho_dk','cang_den','khach','dk_kho','tau','kho','tem','note','ngay_den_kho','nha_xe','so_xe','hien_trang','etd','eta','hang_tau','gia','phu_phi','tong','gio_vao','gio_ra','air','h_bai','h_cang','phi_dien','tong_cp','phat_sinh','cskh','cls']
    rows = []
    for r in range(3, ws.max_row + 1):
        if ws.cell(r, 3).value is None and val(r, 1) is None and val(r, 2) is None: continue
        rec = {'row': r}
        for i, k in enumerate(cols, start=1): rec[k] = val(r, i)
        rows.append(rec)
    return rows

def read_kd(wb):
    ws = wb['Kế hoạch Hạ - Kiểm Dịch']; mm = merged_map(ws)
    val = lambda r, c: mm.get((r, c), ws.cell(r, c).value)
    out = {}
    for r in range(3, ws.max_row + 1):
        left_lo = norm(val(r, 1)); right_lo = norm(val(r, 9))
        rec = {'row': r, 'lo_trai': left_lo, 'etd_trai': d(val(r, 2)), 'cls_eport_trai': ts(val(r, 3)), 'cont_kd': norm(val(r, 4)).upper(), 'seal_kd': norm(val(r, 5)), 'vi_tri_kd': norm(val(r, 6)), 'dia_chi': norm(val(r, 7)),
               'lo_phai': right_lo, 'so_luong': val(r, 10), 'kiem_dich': val(r, 11), 'eport': val(r, 12), 'thanh_ly': val(r, 13), 'etd_phai': d(val(r, 14)), 'cls_eport_phai': ts(val(r, 15)), 'tau': norm(val(r, 16)), 'cls_mail': ts(val(r, 17)), 'so_cont': norm(val(r, 18)), 'cho_keo': val(r, 19), 'kdtv': norm(val(r, 20)), 'note': norm(val(r, 21)), 'to_khai': norm(val(r, 22))}
        lo = re.sub(r'\s*R$', '', right_lo or left_lo).strip()   # '617 R' → 617 (R = rớt / xếp lại)
        if not lo: continue
        rec['lo'] = lo; rec['rot'] = bool(re.search(r'\sR$', right_lo or left_lo))
        out.setdefault(lo, []).append(rec)
    return out

if __name__ == '__main__':
    wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
    rows = read_dieu_do(wb); kd = read_kd(wb)
    lots = OrderedDict()
    for r in rows:
        lo = norm(r['lo'])
        if r['row'] < 1250: continue
        if lo: lots.setdefault(lo, []).append(r)
    for lo, rs in lots.items():
        f = rs[0]
        k = kd.get(lo, [])
        kk = k[-1] if k else {}
        print(lo, '| bk', norm(f['booking']), '| tàu', norm(next((x['tau'] for x in rs if norm(x['tau'])), '')), '| hãng', norm(f['hang_tau']), '| cảng', norm(f['cang_den']), '| ETD', f['etd'] and d(f['etd']), '| ETA', f['eta'] and d(f['eta']), '| CLS AH', norm(f['cls'])[:30], '| cskh', norm(f['cskh']),
              '| KD:', ('ETDt %s clsEt %s KDcont %s | ETDp %s clsEp %s mail %s tàu %s kdtv %s eport %s kd %s tl %s sl %s rot %s' % (kk.get('etd_trai'), kk.get('cls_eport_trai'), kk.get('cont_kd'), kk.get('etd_phai'), kk.get('cls_eport_phai'), kk.get('cls_mail'), kk.get('tau'), kk.get('kdtv'), kk.get('eport'), kk.get('kiem_dich'), kk.get('thanh_ly'), kk.get('so_luong'), kk.get('rot'))) if kk else '-', '| n', len(rs), '| kd rows', len(k))
